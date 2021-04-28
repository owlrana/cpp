import urllib.request
from html_table_parser import HTMLTableParser
import ssl
import pandas as pd
from matplotlib import pyplot as plt
from openpyxl.workbook import Workbook

# Read website content
def url_get_contents(url):

	# Opens a website and read its
	# binary contents (HTTP Response Body)
    context = ssl._create_unverified_context() 
	#making request to the website
    req = urllib.request.Request(url=url)
    f = urllib.request.urlopen(req, context=context)
    return f.read()

# defining the html contents of a URL.
xhtml = url_get_contents('https://www.racefans.net/2021/04/18/2021-emilia-romagna-grand-prix-interactive-data-lap-charts-times-and-tyres/').decode('utf-8')

# Defining the HTMLTableParser object
p = HTMLTableParser()

# feeding the html contents in the HTMLTableParser object
p.feed(xhtml)

# Now finally obtaining the data of the table required
#pprint(p.tables[1])
race_table = p.tables[1]

# for times in timings:
#     print(times)

# print(race_table)

# converting the parsed data to dataframe
print("\n\nPANDAS DATAFRAME\n")
print(pd.DataFrame(p.tables[1]))

timings = list()
names= list()
for i in range(1, len(race_table)):
    time = race_table[i][3].split()
    try:
        timings.append(int(time[0])*60 + float(time[1])) 
        race_table[i][3] = int(time[0])*60 + float(time[1])
    except:
        timings.append(0)


for i in range(1, len(race_table)):
    race_table[i][1] = race_table[i][1][:3].upper()
    names.append(race_table[i][1])
    
    
plt.bar(names,timings,color = 'maroon', width = 0.4) 
plt.xlabel("Drivers")
plt.ylabel("Fastest lap")   
plt.show()


    
    
